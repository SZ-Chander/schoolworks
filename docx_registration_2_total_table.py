import os
from docx import Document
import shutil

dir_path = './docx/ram_docx/' #your docx packet path

def ana_dir(use_dir,list_files):
    re_file_list = []
    doc,docx,mis = 0,0,0
    for file in list_files:
        file_type = file.split('.')[-1]
        if(file_type == 'doc'):
            doc += 1
        elif(file_type == 'docx'):
            docx += 1
            re_file_list.append(new_dir + file)
    print('成功读取到{}份docx文件和{}份doc文件，.doc文件请手动处理'.format(docx,doc))
    return re_file_list

def table_2_data(file_path):
    title = ['队伍名称','队伍口号','队伍成员', '姓名', '性别', '学号', '学院', '专业','负责人联系方式（队长）','队长','队员']
    total_list = []
    document = Document(file_path)
    table = document.tables[0]
    for i in range(len(table.rows)):
        list_result = []
        for j in range(7):
            result = table.cell(j,i).text
            if( result not in title):
                if(result != ''):
                    list_result.append(result)
        if(len(list_result) != 0):
            total_list.append(list_result)
    return total_list

def data_2_line(data):
    team = data[0][0]
    slogan = data[0][1]
    team_people = len(data[0]) - 2
    contrl_key = 0
    team_mess = []
    while(contrl_key < team_people):
        people_mess = []
        for num in range(len(data)-1):
            people_mess.append((data[num][contrl_key + 2]).replace('\n',' ').replace(',','，'))
        contrl_key += 1
        team_mess.append(people_mess)
    team_mess[0].append((data[-1][-1]).replace('\n',' '))
    line = ''
    for n in range(len(team_mess)):
        for j in range(len(team_mess[n])):
            line += (team_mess[n][j]) + ','
    line = team + ',' + slogan + ','+ line
    line += '\n'
    # print(line)
    return line

if __name__ == '__main__':
    title_str = "队伍名称,队伍口号,队长,性别,学号,学院,专业,负责人联系方式（队长),队员,性别,学号,学院,专业\n"
    csv = [title_str]
    error = 0
    for dir in os.listdir(dir_path):
        new_dir = dir_path+dir+'/'
        all_files = ana_dir(new_dir,os.listdir(new_dir))

        for file in all_files:
            data_list = table_2_data(file)
            try:
                new_line = data_2_line(data_list)
            except:
                error += 1
                new_line = '\n'
                print("Error in {},total error {} files".format(file,error))
                goto_dir = './docx/error/'+ str(error)+'/' #input the error file
                os.mkdir(goto_dir)
                shutil.copy(file,goto_dir)
                continue
            csv.append(new_line)

    with open('./docx_2_csv.csv','w') as csv_r:
        for line in csv:
            csv_r.write(line)
